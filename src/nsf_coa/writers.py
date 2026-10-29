"""Output writers: the NSF COA Excel template, a CSV of every row, and a Markdown review report."""

from __future__ import annotations

import copy
import csv
import datetime as dt
import re
import warnings
from importlib import resources
from pathlib import Path

import openpyxl
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

from nsf_coa.merge import COAResult
from nsf_coa.models import Entry

DATE_FORMAT = "mm/dd/yyyy"


def template_path() -> Path:
    return Path(str(resources.files("nsf_coa") / "data" / "coa_template.xlsx"))


# ------------------------------------------------------------------ XLSX ---
def _find_tables(ws) -> dict[int, tuple[int, object]]:
    """Map NSF table number -> (header row, openpyxl Table) by reading column A headers."""
    by_header_row = {CellRange(t.ref).min_row: t for t in ws.tables.values()}
    found = {}
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if value is not None and re.fullmatch(r"[1-5]", str(value).strip()) and row in by_header_row:
            found[int(str(value).strip())] = (row, by_header_row[row])
    return found


def _shift_range(rng: CellRange, at: int, n: int, extend_end_row: int | None) -> CellRange:
    if rng.min_row >= at:
        return CellRange(min_col=rng.min_col, min_row=rng.min_row + n, max_col=rng.max_col, max_row=rng.max_row + n)
    if extend_end_row is not None and rng.max_row == extend_end_row:
        return CellRange(min_col=rng.min_col, min_row=rng.min_row, max_col=rng.max_col, max_row=rng.max_row + n)
    return rng


def _insert_rows(ws, at: int, n: int, table, style_row: int) -> None:
    """Insert ``n`` rows at ``at`` (just below ``table``), keeping tables, validations, and heights intact."""
    last_row = at - 1
    heights = {r: d.height for r, d in ws.row_dimensions.items() if r >= at and d.height}
    ws.insert_rows(at, n)
    for r in sorted(heights, reverse=True):
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        ws.row_dimensions[r + n].height = h

    for t in ws.tables.values():
        rng = CellRange(t.ref)
        new = _shift_range(rng, at, n, last_row if t is table else None)
        t.ref = new.coord
        if t.autoFilter is not None:
            t.autoFilter.ref = new.coord

    for dv in ws.data_validations.dataValidation:
        ranges = [_shift_range(CellRange(r.coord), at, n, last_row) for r in dv.sqref.ranges]
        dv.sqref = MultiCellRange(" ".join(r.coord for r in ranges))

    for r in range(at, at + n):
        for c in range(1, 6):
            ws.cell(r, c)._style = copy.copy(ws.cell(style_row, c)._style)


def _date_cell(cell, value: dt.date | None) -> None:
    if value:
        cell.value = dt.datetime(value.year, value.month, value.day)
        cell.number_format = DATE_FORMAT
    else:
        cell.value = None


def write_xlsx(result: COAResult, path: str | Path, template: str | Path | None = None) -> Path:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # template's Mac-only extension
        wb = openpyxl.load_workbook(template or template_path())
    ws = wb.worksheets[0]
    tables = _find_tables(ws)
    if set(tables) != {1, 2, 3, 4, 5}:
        raise ValueError(f"could not locate COA tables 1-5 in template (found {sorted(tables)})")

    rows_for = {
        1: result.table1 or [Entry(code="1", name="")],
        2: result.table2,
        3: result.table3,
        4: result.table4,
        5: result.table5,
    }

    for number in (5, 4, 3, 2, 1):  # bottom-up so earlier header rows don't move
        header, table = tables[number]
        last = CellRange(table.ref).max_row
        capacity = last - header
        entries = rows_for[number]
        style_row = header + 2 if capacity >= 3 else header + 1
        if len(entries) > capacity:
            _insert_rows(ws, last + 1, len(entries) - capacity, table, style_row)
            last = CellRange(table.ref).max_row
        width = 4 if number in (1, 3) else 5
        for r in range(header + 1, last + 1):  # clear the template's example rows
            for c in range(1, width + 1):
                ws.cell(r, c).value = None
                ws.cell(r, c).hyperlink = None  # else openpyxl writes the link target back as the value
        for i, e in enumerate(entries):
            r = header + 1 + i
            if number == 1:
                ws.cell(r, 2).value = result.self_name if i == 0 else None
                ws.cell(r, 3).value = e.affiliation or None
                _date_cell(ws.cell(r, 4), e.last_active)
                continue
            ws.cell(r, 1).value = f"{e.code}:"
            ws.cell(r, 2).value = e.name
            if number == 2:
                ws.cell(r, 3).value = e.relationship or None
            else:
                ws.cell(r, 3).value = e.affiliation or None
            ws.cell(r, 4).value = e.optional or None
            if number != 3:
                _date_cell(ws.cell(r, 5), e.last_active)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ------------------------------------------------------------------- CSV ---
CSV_FIELDS = ["table", "code", "name", "affiliation", "optional", "relationship", "last_active",
              "orcid", "openalex_ids", "n_works", "most_recent_work", "most_recent_doi", "origin", "flags"]


def _csv_row(table: str, e: Entry) -> dict:
    recent = e.works[0] if e.works else None
    return {
        "table": table, "code": e.code, "name": e.name, "affiliation": e.affiliation,
        "optional": e.optional, "relationship": e.relationship,
        "last_active": e.last_active.strftime("%m/%d/%Y") if e.last_active else "",
        "orcid": e.orcid or "", "openalex_ids": " ".join(sorted(e.openalex_ids)),
        "n_works": len(e.works), "most_recent_work": recent.title if recent else "",
        "most_recent_doi": recent.doi or "" if recent else "", "origin": e.origin,
        "flags": "; ".join(e.flags),
    }


def write_csv(result: COAResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_FIELDS)
        writer.writeheader()
        for e in result.table1:
            writer.writerow(_csv_row("1", e) | {"name": result.self_name})
        for number, entries in ((2, result.table2), (3, result.table3), (4, result.table4), (5, result.table5)):
            for e in entries:
                writer.writerow(_csv_row(str(number), e))
        for e in result.dropped:
            writer.writerow(_csv_row("dropped", e))
    return path


# -------------------------------------------------------------- Markdown ---
def _md(text) -> str:
    return str(text or "").replace("|", "\\|").replace("\n", " ")


def _work_ref(w) -> str:
    label = _md(w.title[:90])
    when = f" ({w.date})" if w.date else ""
    if w.doi:
        return f"[{label}](https://doi.org/{w.doi}){when}"
    return f"{label}{when}"


def write_report(result: COAResult, path: str | Path) -> Path:
    t4a = [e for e in result.table4 if e.code == "A"]
    t4c = [e for e in result.table4 if e.code == "C"]
    flagged = [e for e in result.table4 + result.table5 if e.flags]
    lines = [
        f"# NSF COA review report: {result.self_name}",
        "",
        f"Generated {dt.date.today():%Y-%m-%d} for an as-of date of **{result.as_of:%m/%d/%Y}**.",
        "",
        "Before you upload, check every row against your own records. The data comes from open sources, so",
        "it can be incomplete or wrong, especially for common names and for people who have moved.",
        "",
        "## Summary",
        "",
        "| Table | Rows |",
        "|---|---|",
        f"| 1. Your affiliations (12 mo) | {len(result.table1)} |",
        f"| 2. Relationships (R) | {len(result.table2)} |",
        f"| 3. Ph.D. advisors / advisees (G/T) | {len(result.table3)} |",
        f"| 4. Co-authors (A) | {len(t4a)} |",
        f"| 4. Collaborators (C) | {len(t4c)} |",
        f"| 5. Editorial (B/E) | {len(result.table5)} |",
        f"| Publications used | {len(result.works)} |",
        f"| Rows dropped (out of window / excluded) | {len(result.dropped)} |",
        "",
    ]
    if result.warnings:
        lines += ["## Source warnings", ""] + [f"- {_md(w)}" for w in result.warnings] + [""]
    if flagged:
        lines += ["## Rows that need a look", "", "| Code | Name | Affiliation | Why |", "|---|---|---|---|"]
        lines += [f"| {e.code} | {_md(e.name)} | {_md(e.affiliation)} | {_md('; '.join(e.flags))} |" for e in flagged]
        lines.append("")
    if result.works_without_authors:
        lines += ["## Works with no author list (add co-authors by hand or supply a DOI)", ""]
        lines += [f"- {_work_ref(w)}; source: {', '.join(sorted(w.sources))}" for w in result.works_without_authors]
        lines.append("")
    if result.works_truncated:
        lines += ["## Works with incomplete author lists (\"and others\")", "",
                  "Add the missing co-authors by hand, or add the DOI to `sources.dois` so it can be looked up.", ""]
        lines += [f"- {_work_ref(w)}; listed: {_md(', '.join(a.name for a in w.authors))}"
                  for w in result.works_truncated]
        lines.append("")
    if result.works_self_not_found:
        lines += ["## Works where you were not found in the author list (skipped)",
                  "",
                  "These could be name variants. If any are yours, add the variant to `person.aliases` in your profile.",
                  ""]
        lines += [f"- {_work_ref(w)}; authors: {_md(', '.join(a.name for a in w.authors[:6]))}"
                  f"{' ...' if len(w.authors) > 6 else ''}" for w in result.works_self_not_found]
        lines.append("")
    if result.dropped:
        lines += ["## Dropped rows", "", "| Code | Name / affiliation | Last active | Reason |", "|---|---|---|---|"]
        for e in result.dropped:
            when = e.last_active.strftime("%m/%d/%Y") if e.last_active else ""
            lines.append(f"| {e.code} | {_md(e.name or e.affiliation)} | {when} | {_md('; '.join(e.flags))} |")
        lines.append("")
    lines += ["## Evidence for co-authors and collaborators", "",
              "| Code | Name | Last active | Works | Most recent shared work |", "|---|---|---|---|---|"]
    for e in result.table4:
        when = e.last_active.strftime("%m/%d/%Y") if e.last_active else ""
        recent = _work_ref(e.works[0]) if e.works else _md(e.origin)
        lines.append(f"| {e.code} | {_md(e.name)} | {when} | {len(e.works)} | {recent} |")
    lines.append("")
    if result.works:
        lines += ["## Publications used", ""]
        for w in sorted(result.works, key=lambda w: w.date.end if w.date else dt.date.min, reverse=True):
            lines.append(f"- {_work_ref(w)}; {len(w.authors)} authors; sources: {', '.join(sorted(w.sources))}")
        lines.append("")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path
