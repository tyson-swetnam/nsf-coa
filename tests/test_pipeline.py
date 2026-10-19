import datetime as dt
import warnings

import openpyxl
from conftest import FakeHttp
from openpyxl.worksheet.cell_range import CellRange

from nsf_coa.coa_reader import previous_coa_to_profile_dict, read_previous_coa
from nsf_coa.collect import collect
from nsf_coa.merge import COAResult, build_coa
from nsf_coa.models import Entry
from nsf_coa.writers import template_path, write_csv, write_report, write_xlsx


def _http(load_json):
    return FakeHttp({
        "https://api.openalex.org/works": load_json("openalex_works.json"),
        "https://api.crossref.org/works": {"message": {"items": []}},
        "https://pub.orcid.org/v3.0/0000-0001-0000-0001/works": {"group": []},
    })


def _previous_coa(path):
    prev = COAResult(self_name="Researcher, Rita Q", as_of=dt.date(2024, 1, 1))
    prev.table1 = [Entry("1", "", affiliation="University of Example"),
                   Entry("1", "", affiliation="Old University", last_active=dt.date(2025, 1, 1))]
    prev.table2 = [Entry("R", "Spouse, Sam", relationship="Spouse")]
    prev.table3 = [Entry("T", "Student, Stu", affiliation="Agency X", optional="stu@example.org")]
    prev.table4 = [Entry("A", "Old, Friend", affiliation="Somewhere", last_active=dt.date(2020, 1, 1)),
                   Entry("A", "Kept, Karl", affiliation="Karl Inst", last_active=dt.date(2024, 5, 1)),
                   Entry("A", "Builder, Bob", affiliation="Old College", optional="bob@new.edu"),
                   Entry("C", "Granted, Gail", affiliation="Grant U")]
    prev.table5 = [Entry("B", "Editor, Ed", optional="J Old", last_active=dt.date(2023, 1, 1)),
                   Entry("E", "Coeditor, Cy", optional="J New", last_active=dt.date(2025, 5, 1))]
    return write_xlsx(prev, path)


def test_end_to_end(tmp_path, fixtures, load_json, profile):
    prev_path = _previous_coa(tmp_path / "prev.xlsx")
    prof = profile(
        sources={"bibtex": [str(fixtures / "scholar.bib")], "datacite": False},
        collaborators=[{"name": "Dana Partner", "affiliation": "Partner Lab", "last_active": "2026-01-01"}],
        previous_coa=str(prev_path),
    )
    collection = collect(prof, _http(load_json), progress=lambda *_: None)
    result = build_coa(prof, collection.works, read_previous_coa(prev_path))

    names = {(e.code, e.name) for e in result.table4}
    assert ("A", "Builder, Bob") in names  # "B. Builder" merged into "Bob Builder"
    assert ("A", "de la Cruz, Carla") in names
    assert ("A", "Müller, Jörg") in names and ("A", "Zhang, Wei") in names
    assert ("C", "Partner, Dana") in names and ("A", "Partner, Dana") not in names  # listed once, under C
    assert ("A", "Kept, Karl") in names  # carried forward from previous COA
    assert ("C", "Granted, Gail") in names
    assert not any("Researcher" in n for _, n in names)  # never list yourself
    assert not any(n in {"Errata Person", "Person, Errata", "Friend, Old"} for _, n in names)

    bob = next(e for e in result.table4 if e.name == "Builder, Bob")
    assert bob.optional == "bob@new.edu"  # curated email from previous COA kept
    assert bob.last_active == dt.date(2025, 3, 14) and len(bob.works) == 2

    dropped = {e.name or e.affiliation for e in result.dropped}
    assert {"Old, Friend", "Editor, Ed", "Old University"} <= dropped
    assert [e.name for e in result.table2] == ["Spouse, Sam"]
    assert [(e.code, e.name) for e in result.table3] == [("T", "Student, Stu")]
    assert [e.name for e in result.table5] == ["Coeditor, Cy"]
    assert [e.affiliation for e in result.table1] == ["University of Example"]
    assert [w.title for w in result.works_truncated] == ["Open science workshops in practice"]

    out = write_xlsx(result, tmp_path / "coa.xlsx")
    write_csv(result, tmp_path / "coa.csv")
    report = write_report(result, tmp_path / "report.md").read_text()
    assert "Kept, Karl" in report and "Dropped rows" in report

    # Round trip: what we wrote reads back as the same rows.
    again = read_previous_coa(out)
    assert {(r["code"], r["name"]) for r in again.rows} == {
        (e.code, e.name) for e in result.table2 + result.table3 + result.table4 + result.table5}
    assert again.self_name == "Researcher, Rita Q"


def test_xlsx_grows_tables_and_keeps_structure(tmp_path):
    result = COAResult(self_name="Doe, Jane", as_of=dt.date(2026, 1, 1))
    result.table1 = [Entry("1", "", affiliation=f"Org {i}") for i in range(7)]
    result.table4 = [Entry("A", f"Person{i:03d}, Pat", affiliation="Uni", last_active=dt.date(2025, 1, 1))
                     for i in range(120)]
    result.table5 = [Entry("E", "Co, Ed", optional="Journal")]
    path = write_xlsx(result, tmp_path / "big.xlsx")

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(path).active
        template = openpyxl.load_workbook(template_path()).active
    refs = {t.name: CellRange(t.ref) for t in ws.tables.values()}
    old = {t.name: CellRange(t.ref) for t in template.tables.values()}
    assert refs["TableA"].max_row - refs["TableA"].min_row == 7
    assert refs["TableD"].max_row - refs["TableD"].min_row == 120
    shift = 2 + (120 - 5)
    assert refs["TableD5"].min_row == old["TableD5"].min_row + shift
    # Header text still matches the table column names.
    assert ws.cell(refs["TableD"].min_row, 1).value == "4"
    assert ws.cell(refs["TableD5"].min_row, 1).value == "5"
    # Dropdown validation covers every new Table 4 row.
    dv = next(d for d in ws.data_validations.dataValidation if d.formula1 == '"A:,C:"')
    assert any(r.max_row == refs["TableD"].max_row for r in dv.sqref.ranges)
    last = refs["TableD"].max_row
    assert ws.cell(last, 2).value == "Person119, Pat"
    assert ws.cell(last, 5).number_format == "mm/dd/yyyy"
    # Template example rows are gone.
    values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert not any("Alphaman" in v for v in values)


def test_template_examples_are_ignored_by_reader():
    prev = read_previous_coa(template_path())
    assert prev.rows == [] and prev.self_name == ""


def test_import_coa_yaml(tmp_path):
    prev_path = _previous_coa(tmp_path / "prev.xlsx")
    data = previous_coa_to_profile_dict(read_previous_coa(prev_path))
    assert data["relationships"][0]["relationship"] == "Spouse"
    assert data["advisors"][0] == {"type": "T", "name": "Student, Stu", "affiliation": "Agency X",
                                   "optional": "stu@example.org"}
    assert data["editorial"][1]["journal"] == "J New"


def test_exclude_and_override(tmp_path, load_json, profile):
    prof = profile(
        exclude=["de la Cruz, Carla"],
        overrides={"Builder, Bob": {"affiliation": "New Employer", "optional": "bob@x.org"}},
        sources={"datacite": False},
    )
    collection = collect(prof, _http(load_json), progress=lambda *_: None)
    result = build_coa(prof, collection.works)
    names = {e.name: e for e in result.table4}
    assert "de la Cruz, Carla" not in names
    assert names["Builder, Bob"].affiliation == "New Employer"


def test_initials_ambiguity_is_flagged(profile):
    from nsf_coa.models import AuthorMention, PubDate, Work

    me = AuthorMention("Rita Q. Researcher", orcid="0000-0001-0000-0001")
    works = [
        Work("one", PubDate(2025, 1, 1), [me, AuthorMention("John Smith")], doi="10.1111/1"),
        Work("two", PubDate(2025, 2, 1), [me, AuthorMention("Jane Smith")], doi="10.1111/2"),
        Work("three", PubDate(2025, 3, 1), [me, AuthorMention("J. Smith")], doi="10.1111/3"),
    ]
    result = build_coa(profile(), works)
    j = next(e for e in result.table4 if e.name == "Smith, J")
    assert any("ambiguous" in f for f in j.flags)
    assert {e.name for e in result.table4} == {"Smith, John", "Smith, Jane", "Smith, J"}
